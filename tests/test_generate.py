import unittest

from pptx import Presentation
from pptx.oxml.ns import qn

from app import config
from app.routers.generate import generation_suffix
from app.services.builder import append_slide_from_source, clone_slide


class GenerateNamingTests(unittest.TestCase):
    def test_editor_download_is_named_as_edited_copy(self):
        self.assertEqual(generation_suffix({}, edited_copy=True), "_수정본")

    def test_saved_layout_uses_edited_suffix_for_legacy_generate(self):
        overrides = {"global": {"title": {"font_size": 22}}, "slides": {}}
        self.assertEqual(generation_suffix(overrides), "_편집본")

    def test_plain_generate_keeps_auto_layout_suffix(self):
        self.assertEqual(generation_suffix({}), "_자동배열본")


class PowerPointSlideStructureTests(unittest.TestCase):
    def assert_single_root_group_nodes(self, slide):
        tags = [child.tag for child in slide.shapes._spTree]
        self.assertEqual(tags.count(qn("p:nvGrpSpPr")), 1)
        self.assertEqual(tags.count(qn("p:grpSpPr")), 1)

    def test_template_clone_copies_only_shapes(self):
        prs = Presentation(config.DEFAULT_TEMPLATE)
        cloned = clone_slide(prs, 0)
        self.assert_single_root_group_nodes(cloned)

    def test_merge_append_copies_only_shapes(self):
        source = Presentation(config.DEFAULT_TEMPLATE)
        dest = Presentation(config.DEFAULT_TEMPLATE)
        append_slide_from_source(dest, source.slides[0])
        self.assert_single_root_group_nodes(dest.slides[-1])


if __name__ == "__main__":
    unittest.main()
