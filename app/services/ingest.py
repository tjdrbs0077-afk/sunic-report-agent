"""PPTX 추출 — _ref/report_ai_prototype/ppt_ingest.py 이식."""
from __future__ import annotations

import re
import uuid
from collections import Counter
from pathlib import Path
from typing import Any, Iterable

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

EMU = 914400
DATE_RE = re.compile(r"(?:['’]?\d{2}[.\-/](?:Q?[1-4]|\d{1,2})(?:월)?|\d{4}[.\-/]\d{1,2}(?:[.\-/]\d{1,2})?|Q[1-4])", re.I)
STOPWORDS = {
    "그리고", "또한", "관련", "대한", "위한", "통해", "사업", "추진", "내용", "보고", "검토", "현황",
    "계획", "주요", "기반", "중심", "대상", "구분", "페이지", "자료", "결과", "예정", "필요",
}

# 원본 슬라이드에 보이는 섹션 코드·번호·서식 가이드 텍스트를 실제 보고 내용과 분리한다.
SECTION_CODE_RE = re.compile(r"^\d{1,2}\.\s*[A-Z0-9 &/\-]+$")
_CIRCLED_NUMBERS = "①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳"
_CIRCLED_NUMBER_RE = re.compile(rf"^[{_CIRCLED_NUMBERS}]$")
_MARKER_ONLY_RE = re.compile(
    rf"^(?:[{_CIRCLED_NUMBERS}]|\d{{1,2}}[.)]?|[▶▷◆◇■□※✓√●○◎☞▸»*·◦▪❑–—•])$"
)
_NUMBER_PREFIX_RE = re.compile(r"^(?:\d{1,2}[.)]|\(\d{1,2}\))\s+")
_CIRCLED_PREFIX_RE = re.compile(rf"^[{_CIRCLED_NUMBERS}]\s*")
_BULLET_PREFIX_RE = re.compile(r"^[▶▷◆◇■□※✓√●○◎☞▸»*·◦▪❑–—•]\s*")
_NUMERIC_METRIC_RE = re.compile(r"^[\d\s.,%+:/\-]+$")


def _is_format_guide_text(text: str) -> bool:
    """템플릿 제작자가 남긴 글꼴 안내 말풍선을 내용에서 제외한다."""
    compact = re.sub(r"[\s+()/_\-]", "", text or "").casefold()
    if not compact:
        return False
    token = r"(?:나눔스퀘어(?:otf|ac)?|맑은고딕|corbel|arial|tahoma|wingdings|extrabold|bold|볼드)"
    return re.fullmatch(rf"(?:{token})+", compact, re.I) is not None


def _is_marker_only(text: str) -> bool:
    return _MARKER_ONLY_RE.fullmatch(clean_text(text)) is not None


def strip_explicit_list_prefix(text: str) -> str:
    """이미 들어간 1./1)/①/불릿을 제거해 표준 번호가 중복되지 않게 한다."""
    value = clean_text(text)
    if not value:
        return ""
    if _is_marker_only(value):
        return ""
    for pattern in (_NUMBER_PREFIX_RE, _CIRCLED_PREFIX_RE, _BULLET_PREFIX_RE):
        replaced = pattern.sub("", value, count=1).strip()
        if replaced != value:
            return replaced
    return value


def normalize_body_items(items: list[dict[str, Any]], page_title: str) -> list[dict[str, Any]]:
    """추출·저장 시점과 무관하게 본문을 표준 번호를 붙일 수 있는 상태로 만든다."""
    title_key = clean_text(page_title).casefold()
    cleaned: list[dict[str, Any]] = []
    seen: set[str] = set()
    for item in items or []:
        raw = clean_text(str(item.get("text", "")))
        if (
            not raw
            or raw.casefold() == title_key
            or SECTION_CODE_RE.fullmatch(raw)
            or _is_format_guide_text(raw)
            or _is_marker_only(raw)
        ):
            continue
        text = strip_explicit_list_prefix(raw)
        key = text.casefold()
        if not text or key == title_key or key in seen:
            continue
        seen.add(key)
        copied = dict(item)
        copied["text"] = text
        copied["level"] = max(0, min(4, int(item.get("level", 0) or 0)))
        copied["bold"] = item.get("bold")
        cleaned.append(copied)

    if not cleaned:
        return []

    # 제목은 별도 상자에 있으므로 본문 최상위는 항상 1단계부터 시작한다.
    min_level = min(int(item["level"]) for item in cleaned)
    if min_level:
        for item in cleaned:
            item["level"] = int(item["level"]) - min_level
    previous = 0
    for index, item in enumerate(cleaned):
        level = int(item["level"])
        if index == 0:
            level = 0
        else:
            level = min(level, previous + 1)
        item["level"] = level
        previous = level
    return cleaned


def normalize_payload_content(payload: dict[str, Any]) -> dict[str, Any]:
    """기존 JSON도 재업로드 없이 새 정규화 규칙을 적용한다."""
    for key in ("slides", "original_slides"):
        for slide in payload.get(key, []) or []:
            slide["body"] = normalize_body_items(slide.get("body", []), slide.get("page_title", ""))
    return payload


def inch(value: int | float) -> float:
    return float(value) / EMU


def clean_text(value: str) -> str:
    return re.sub(r"\s+", " ", (value or "").replace("\x0b", " ")).strip()


def safe_stem(filename: str) -> str:
    stem = Path(filename).stem.strip() or "업로드 보고서"
    return re.sub(r"[\\/:*?\"<>|]+", "_", stem)[:80]


def _iter_shapes(shapes) -> Iterable[Any]:
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)


def _font_size_pt(paragraph) -> float:
    sizes: list[float] = []
    for run in paragraph.runs:
        if run.font.size:
            sizes.append(float(run.font.size.pt))
    return max(sizes) if sizes else 11.0


def _shape_text_entries(slide, slide_h: float) -> list[dict[str, Any]]:
    entries: list[dict[str, Any]] = []
    for shape in _iter_shapes(slide.shapes):
        if shape.shape_type == MSO_SHAPE_TYPE.TABLE or not getattr(shape, "has_text_frame", False):
            continue
        x, y, w, h = map(inch, (shape.left, shape.top, shape.width, shape.height))
        for pi, paragraph in enumerate(shape.text_frame.paragraphs):
            text = clean_text(paragraph.text)
            if not text:
                continue
            entries.append({
                "text": text,
                "x": x,
                "y": y,
                "w": w,
                "h": h,
                "bottom": y + h,
                "level": int(getattr(paragraph, "level", 0) or 0),
                "font_size": _font_size_pt(paragraph),
                "bold": any(bool(run.font.bold) for run in paragraph.runs),
                "paragraph_index": pi,
                "shape_name": shape.name,
                "is_placeholder": bool(getattr(shape, "is_placeholder", False)),
                "is_bottom": y > slide_h * 0.84,
            })
    return entries


def _extract_tables(slide) -> list[dict[str, Any]]:
    tables: list[dict[str, Any]] = []
    for shape in _iter_shapes(slide.shapes):
        if shape.shape_type != MSO_SHAPE_TYPE.TABLE:
            continue
        table = shape.table
        matrix = [[clean_text(table.cell(r, c).text) for c in range(len(table.columns))] for r in range(len(table.rows))]
        if not matrix:
            continue
        headers = matrix[0]
        rows = [row for row in matrix[1:] if any(row)]
        tables.append({"headers": headers, "rows": rows, "y": inch(shape.top)})
    return sorted(tables, key=lambda x: x["y"])


def _choose_title(entries: list[dict[str, Any]], slide_h: float) -> dict[str, Any] | None:
    candidates = [
        e for e in entries
        if not e["is_bottom"]
        and not _is_format_guide_text(e["text"])
        and not _is_marker_only(e["text"])
        and not _NUMERIC_METRIC_RE.fullmatch(e["text"])
        and not SECTION_CODE_RE.fullmatch(e["text"])
    ]
    # 일반 내용 페이지는 최상단 제목 띠를 우선한다. 표지처럼 제목이 중앙에 있는
    # 슬라이드는 최상단 후보가 없으므로 전체 후보로 자연스럽게 폴백한다.
    top_band = [e for e in candidates if e["y"] < min(1.45, slide_h * 0.22)]
    if top_band:
        candidates = top_band
    if not candidates:
        return None
    # 위쪽 위치·큰 폰트·placeholder·짧은 제목형 문자열에 가중치 합산.
    # (v4는 튜플 비교라 placeholder 여부가 폰트 크기를 무조건 이겨 24pt 제목 상자가 밀리는 문제가 있었음)
    return max(
        candidates,
        key=lambda e: (
            (2 if e["is_placeholder"] else 0)
            + e["font_size"] * 2.0
            + max(0, 5 - e["y"] * 2)
            + (3 if e["x"] < 1.5 else 0)
            + min(e["w"], 8) * 0.25
            + (2 if 4 <= len(e["text"]) <= 70 else 0)
            - (5 if len(e["text"]) <= 2 else 0)
        ),
    )


def _estimate_level(entry: dict[str, Any], max_font: float) -> int:
    if entry["level"]:
        return max(0, min(4, entry["level"]))
    delta = max_font - entry["font_size"]
    level = 0 if delta < 1 else 1 if delta < 2.5 else 2 if delta < 4.5 else 3
    return max(0, min(4, level))


def _circled_number(text: str) -> int | None:
    value = clean_text(text)
    return _CIRCLED_NUMBERS.index(value) + 1 if _CIRCLED_NUMBER_RE.fullmatch(value) else None


def _order_body_entries(entries: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """①~④ 카드형 레이아웃은 숫자 순서대로 제목과 설명을 묶어 읽는다."""
    markers = [(entry, _circled_number(entry["text"])) for entry in entries]
    markers = [(entry, number) for entry, number in markers if number is not None]
    content = [entry for entry in entries if _circled_number(entry["text"]) is None]
    if len(markers) < 2:
        return sorted(content, key=lambda e: (round(e["y"], 2), round(e["x"], 2), e["paragraph_index"]))

    column_xs: list[float] = []
    for marker, _ in sorted(markers, key=lambda pair: pair[0]["x"]):
        if not column_xs or abs(marker["x"] - column_xs[-1]) > 1.5:
            column_xs.append(marker["x"])
    bounds = [-float("inf")]
    bounds.extend((left + right) / 2 for left, right in zip(column_xs, column_xs[1:]))
    bounds.append(float("inf"))

    ordered: list[dict[str, Any]] = []
    used: set[int] = set()
    for marker, _ in sorted(markers, key=lambda pair: pair[1]):
        column = min(range(len(column_xs)), key=lambda i: abs(marker["x"] - column_xs[i]))
        next_rows = [
            other["y"] for other, _ in markers
            if abs(other["x"] - marker["x"]) < 1.5 and other["y"] > marker["y"] + 0.2
        ]
        row_end = min(next_rows) - 0.2 if next_rows else float("inf")
        group = [
            entry for entry in content
            if id(entry) not in used
            and bounds[column] <= entry["x"] < bounds[column + 1]
            and marker["y"] - 0.25 <= entry["y"] < row_end
        ]
        group.sort(key=lambda e: (round(e["y"], 2), round(e["x"], 2), e["paragraph_index"]))
        ordered.extend(group)
        used.update(id(entry) for entry in group)

    leftovers = [entry for entry in content if id(entry) not in used]
    leftovers.sort(key=lambda e: (round(e["y"], 2), round(e["x"], 2), e["paragraph_index"]))
    ordered.extend(leftovers)
    return ordered


def _shorten(text: str, limit: int = 180) -> str:
    return text if len(text) <= limit else text[: limit - 1].rstrip() + "…"


def _keywords(texts: list[str], top_k: int = 8) -> list[str]:
    tokens: list[str] = []
    for text in texts:
        for token in re.findall(r"[A-Za-z][A-Za-z0-9+._-]{2,}|[가-힣]{2,}", text):
            t = token.strip("._-")
            if len(t) < 2 or t in STOPWORDS or t.isdigit():
                continue
            tokens.append(t)
    return [token for token, _ in Counter(tokens).most_common(top_k)]


def _normalize_table(spec: dict[str, Any], max_cols: int, max_rows: int) -> dict[str, Any]:
    headers = list(spec.get("headers", []))[:max_cols]
    if not headers:
        headers = ["구분", "주요 내용"][:max_cols]
    rows = []
    for row in spec.get("rows", [])[:max_rows]:
        row2 = list(row)[:max_cols]
        row2 += [""] * (len(headers) - len(row2))
        rows.append(row2)
    while len(rows) < min(2, max_rows):
        rows.append([""] * len(headers))
    return {"headers": headers, "rows": rows}


def extract_presentation(path: Path, report_id: str | None = None, display_name: str | None = None) -> dict[str, Any]:
    prs = Presentation(path)
    if not prs.slides:
        raise ValueError("슬라이드가 없는 PPT입니다.")
    slide_h = inch(prs.slide_height)
    name = display_name or safe_stem(path.name)
    rid = report_id or f"upload_{uuid.uuid4().hex[:10]}"
    slides: list[dict[str, Any]] = []
    all_texts: list[str] = []

    for slide_no, slide in enumerate(prs.slides, 1):
        entries = _shape_text_entries(slide, slide_h)
        tables = _extract_tables(slide)
        title_entry = _choose_title(entries, slide_h)
        page_title = _shorten(title_entry["text"], 70) if title_entry else f"페이지 {slide_no}"

        foot_entries = [e for e in entries if e["is_bottom"] and e is not title_entry and not re.fullmatch(r"[‹<>#\s\d/]+", e["text"])]
        footnote = _shorten(" · ".join(e["text"] for e in foot_entries[:2]), 220)
        body_entries = [
            e for e in entries
            if e is not title_entry and e not in foot_entries
            and not re.fullmatch(r"[‹<>#\s\d/]+", e["text"])
            and e["text"] not in {"구분", "구 분", "주요 내용", "Project 진행 일정(案)"}
            # 표준 양식 장식 텍스트(표 캡션 【 표 】 등)는 본문이 아니다
            and not re.fullmatch(r"【.*】", e["text"].replace(" ", ""))
            and not _is_format_guide_text(e["text"])
        ]
        body_entries = _order_body_entries(body_entries)
        semantic_entries = [e for e in body_entries if not _is_marker_only(e["text"])]
        max_font = max((e["font_size"] for e in semantic_entries), default=12)
        body: list[dict[str, Any]] = []
        seen: set[str] = set()
        for entry in body_entries:
            text = _shorten(entry["text"], 180)
            key = text.casefold()
            if not text or key in seen:
                continue
            # 페이지 제목과 같은 줄, "08. FINANCIAL OUTLOOK" 같은 섹션 코드는
            # 본문이 아니라 장식이다 — 표준 양식 본문에 포함하지 않는다.
            if key == page_title.strip().casefold() or SECTION_CODE_RE.fullmatch(text):
                continue
            if _is_marker_only(text):
                continue
            seen.add(key)
            body.append({"text": text, "level": _estimate_level(entry, max_font), "bold": bool(entry["bold"]) or None})
            if len(body) >= 30:
                break
        body = normalize_body_items(body, page_title)
        # 표지·간지처럼 본문이 없는 장에 제목을 복제하지 않는다. 생성기는 빈 본문을 지원한다.

        date_tokens: list[str] = []
        for entry in entries:
            date_tokens.extend(DATE_RE.findall(entry["text"]))
        date_tokens = list(dict.fromkeys(date_tokens))[:6]
        if len(tables) >= 2:
            template_type = 3
        elif tables:
            template_type = 1
        else:
            template_type = 2

        key_quotes = [page_title] + [item["text"] for item in body[:3]]
        evidence = [
            {"label": "페이지 제목", "quote": page_title, "location": "원본 슬라이드 상단 제목"},
        ]
        for i, quote in enumerate(key_quotes[1:3], 1):
            evidence.append({"label": f"핵심 근거 {i}", "quote": quote, "location": "원본 슬라이드 본문"})
        summary_parts = [item["text"] for item in body[:3] if item["text"] != page_title]
        summary = f"‘{page_title}’ 페이지는 " + ("; ".join(summary_parts) if summary_parts else "원본 슬라이드의 핵심 내용을 정리한다") + "."
        slide_data: dict[str, Any] = {
            "slide_number": slide_no,
            "template_type": template_type,
            "page_title": page_title,
            "sidebar": _shorten(name, 20).replace(" ", "\n", 1),
            "body": body,
            "summary": summary,
            "evidence": evidence,
            "footnote": footnote or f"※ 출처: 업로드 파일 {path.name}, 원본 p.{slide_no}",
            "source_slide_number": slide_no,
        }
        if template_type == 1:
            slide_data["table1"] = _normalize_table(tables[0], 2, 5)
            slide_data["timeline"] = date_tokens + [""] * (6 - len(date_tokens))
            slide_data["timeline_note"] = ["원본 일정"] * len(date_tokens) + [""] * (6 - len(date_tokens))
        elif template_type == 3:
            slide_data["table1"] = _normalize_table(tables[0], 3, 3)
            slide_data["table2"] = _normalize_table(tables[1], 5, 3)
        slides.append(slide_data)
        all_texts.extend([page_title] + [item["text"] for item in body])

    kws = _keywords(all_texts)
    topic = slides[0]["page_title"] if slides else name
    unit = {
        "id": rid,
        "name": name,
        "short": _shorten(name, 20),
        "topic": topic,
        "keywords": kws,
        "partners": [],
        "goal": topic,
        "source": "uploaded",
        "source_file": path.name,
    }
    return {"unit": unit, "slides": slides}
