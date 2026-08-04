"""양식 검증기 — 업로드된 PPTX를 config/standard_rules.yaml 규칙과 대조해 위반 목록을 만든다.

각 위반 항목은 SUNIC UI의 .fixlist / .fb-item 마크업에 그대로 들어가는 형태:
{category, severity, slide_no, message, detail, auto_fixable}
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Any

import yaml
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.util import Inches

from app import config

CATEGORY_SEVERITY = {
    "글머리 기호": "high",
    "글꼴": "high",
    "글자 크기": "mid",
    "들여쓰기": "mid",
    "표 규격": "mid",
    "정리": "low",
}

# 표준 글머리(자동번호 1. / 1) / ❑ / – / •) 외에 텍스트 맨 앞에 직접 쓰인 기호
BAD_BULLET_CHARS = set("▶▷◆◇■□※✓√●○◎☞▸»*·◦▪")

# 표준 3종 폰트 외에 스펙이 허용하는 보조 폰트와 테마 참조
FONT_EXTRAS = {"Tahoma", "Arial", "Wingdings", "+mn-lt", "+mn-ea", "+mj-lt", "+mj-ea", "+mn-cs", "+mj-cs"}

SIZE_TOL = 0.5   # pt
GEOM_TOL = 0.10  # inch
MAX_ISSUES = 500


def load_rules() -> dict[str, Any]:
    if not config.RULES_FILE.exists():
        return {}
    return yaml.safe_load(config.RULES_FILE.read_text(encoding="utf-8")) or {}


def _iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)


def _run_ea_font(run) -> str | None:
    rpr = run._r.find(qn("a:rPr"))
    if rpr is None:
        return None
    ea = rpr.find(qn("a:ea"))
    return ea.get("typeface") if ea is not None else None


def _nearest(value: float, candidates: list[float]) -> float:
    return min(candidates, key=lambda c: abs(c - value))


def validate_pptx(path: Path) -> dict[str, Any]:
    rules = load_rules()
    fonts = rules.get("fonts") or {}
    allowed_fonts = {fonts.get("latin"), fonts.get("korean"), fonts.get("heading_korean")} - {None} | FONT_EXTRAS
    body_levels = rules.get("body_levels") or []
    level_sizes = [float(x.get("size", s)) for x, s in zip(body_levels, [14, 13, 12, 11, 10])] or [14, 13, 12, 11, 10]
    title_size = float((rules.get("title") or {}).get("font_size", 24))
    table_rules = rules.get("table") or {}
    footnote_size = float((rules.get("footnote") or {}).get("font_size", 8))
    cleanup = rules.get("cleanup") or {}
    max_level = len(level_sizes) - 1

    issues: list[dict[str, Any]] = []
    seen: set[tuple] = set()

    def add(category: str, slide_no: int, message: str, detail: str) -> None:
        key = (category, slide_no, message)
        if key in seen or len(issues) >= MAX_ISSUES:
            return
        seen.add(key)
        issues.append({
            "category": category,
            "severity": CATEGORY_SEVERITY[category],
            "slide_no": slide_no,
            "message": message,
            "detail": detail,
            "auto_fixable": True,
        })

    prs = Presentation(path)
    slide_h = float(prs.slide_height) / 914400

    for slide_no, slide in enumerate(prs.slides, 1):
        for shape in _iter_shapes(slide.shapes):
            # 표 규격: 좌측 기준 x_in, 폭 width_in 고정
            if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
                x_in = float(shape.left) / 914400
                w_in = float(shape.width) / 914400
                std_x = float(table_rules.get("x_in", 2.89757))
                std_w = float(table_rules.get("width_in", 7.519097))
                if abs(x_in - std_x) > GEOM_TOL or abs(w_in - std_w) > GEOM_TOL:
                    add("표 규격", slide_no,
                        f"표 위치·폭 조정 필요 (L {x_in:.2f}in → {std_x:.2f}in, W {w_in:.2f}in → {std_w:.2f}in)",
                        f"슬라이드 {slide_no}의 표 — 표준: 좌 7.36cm, 폭 19.10cm 고정")
                continue
            if not getattr(shape, "has_text_frame", False):
                continue

            top_in = float(shape.top) / 914400 if shape.top is not None else 0.0
            left_in = float(shape.left) / 914400 if shape.left is not None else 0.0
            is_title_area = top_in < 0.8 and left_in < 1.2
            # 타임라인 영역(스펙 3-6): 제목 10pt·마디 날짜 9pt·하단 설명 8pt가 표준
            is_timeline_band = 6.1 < top_in < 7.15 and left_in > 2.4
            is_bottom_area = top_in > slide_h * 0.84 and not is_timeline_band

            if not shape.text_frame.text.strip():
                # 도형(선·화살표·마디 원 등)은 텍스트프레임이 비어 있는 것이 정상 — 순수 텍스트상자만 위반
                if cleanup.get("remove_empty_textbox", True) and shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
                    add("정리", slide_no, "빈 텍스트상자 제거 대상", f"슬라이드 {slide_no}의 '{shape.name}'")
                continue

            for paragraph in shape.text_frame.paragraphs:
                text = paragraph.text.strip()
                if not text:
                    continue
                level = int(paragraph.level or 0)
                # 표 캡션 【 표 】 — 스펙 3-4: Corbel 11pt Bold가 표준
                is_caption = text.startswith("【") and text.endswith("】")

                # 글머리 기호 — 하단(각주) 영역의 ※ 표기는 스펙상 정상이므로 제외
                if not is_bottom_area and text[0] in BAD_BULLET_CHARS:
                    add("글머리 기호", slide_no,
                        f"'{text[0]}' → 표준 글머리(1. / 1) / ❑ / – / •) 교체 필요",
                        f"슬라이드 {slide_no}의 {level + 1}단계 문단 “{text[:30]}”")

                # 들여쓰기 — 5단계 초과
                if level > max_level:
                    add("들여쓰기", slide_no,
                        f"{level + 1}단계 문단 → {max_level + 1}단계 이내로 정리 필요",
                        f"슬라이드 {slide_no} “{text[:30]}”")

                # 정리 — 연속 공백
                if cleanup.get("collapse_spaces", True) and re.search(r"  +", paragraph.text):
                    add("정리", slide_no, "연속 공백 → 1칸으로 축소 필요", f"슬라이드 {slide_no} “{text[:30]}”")

                for run in paragraph.runs:
                    if not run.text.strip():
                        continue
                    # 글꼴
                    for font_name in (run.font.name, _run_ea_font(run)):
                        if font_name and font_name not in allowed_fonts:
                            add("글꼴", slide_no,
                                f"'{font_name}' → Corbel · 나눔스퀘어 계열로 교체 필요",
                                f"슬라이드 {slide_no} “{run.text[:30]}”")
                            break
                    # 글자 크기
                    if run.font.size is None:
                        continue
                    pt = float(run.font.size.pt)
                    if is_caption:
                        if abs(pt - 11.0) > SIZE_TOL:
                            add("글자 크기", slide_no,
                                f"표 캡션 {pt:g}pt → 11pt 통일 필요",
                                f"슬라이드 {slide_no}의 표 캡션")
                    elif is_title_area:
                        if abs(pt - title_size) > SIZE_TOL:
                            add("글자 크기", slide_no,
                                f"제목 {pt:g}pt → {title_size:g}pt 통일 필요",
                                f"슬라이드 {slide_no}의 페이지 제목")
                    elif is_timeline_band:
                        if not any(abs(pt - allowed) <= SIZE_TOL for allowed in (8, 9, 10)):
                            add("글자 크기", slide_no,
                                f"타임라인 {pt:g}pt → 9pt(날짜)·8pt(설명) 통일 필요",
                                f"슬라이드 {slide_no}의 타임라인 영역")
                    elif is_bottom_area:
                        if abs(pt - footnote_size) > SIZE_TOL:
                            add("글자 크기", slide_no,
                                f"각주 {pt:g}pt → {footnote_size:g}pt 통일 필요",
                                f"슬라이드 {slide_no}의 하단 각주")
                    else:
                        expected = level_sizes[min(level, max_level)]
                        if abs(pt - expected) > SIZE_TOL:
                            add("글자 크기", slide_no,
                                f"본문 {pt:g}pt → {expected:g}pt 통일 필요",
                                f"슬라이드 {slide_no}의 {level + 1}단계 문단")

    return summarize(issues)


def summarize(issues: list[dict[str, Any]]) -> dict[str, Any]:
    counts: dict[str, int] = {}
    for issue in issues:
        counts[issue["category"]] = counts.get(issue["category"], 0) + 1
    by_category = [
        {"category": category, "count": count, "severity": CATEGORY_SEVERITY[category]}
        for category, count in sorted(counts.items(), key=lambda kv: -kv[1])
    ]
    return {"total": len(issues), "by_category": by_category, "issues": issues}


# ── 규칙 파일 저장 · PPTX에서 규칙 추출 ─────────────────────────

RULE_ORDER = ["canvas", "fonts", "title", "body_levels", "table", "frame", "footnote", "continuation", "cleanup"]


def save_rules(rules: dict[str, Any]) -> dict[str, Any]:
    """규칙을 yaml로 저장한다. 키 순서를 고정해 사람이 읽기 좋게 유지."""
    ordered = {k: rules[k] for k in RULE_ORDER if k in rules}
    ordered.update({k: v for k, v in rules.items() if k not in ordered})
    config.RULES_FILE.parent.mkdir(parents=True, exist_ok=True)
    config.RULES_FILE.write_text(
        yaml.safe_dump(ordered, allow_unicode=True, sort_keys=False, default_flow_style=False),
        encoding="utf-8",
    )
    return ordered


def derive_rules_from_pptx(path: Path) -> dict[str, Any]:
    """기준 양식 PPTX에서 규칙 값을 실측해 standard_rules.yaml 내용을 만든다.

    YAML을 손으로 고치는 대신 새 양식 PPTX를 올리면 기준이 갱신된다.
    """
    from app.services.builder import exact_template_profile

    profile = exact_template_profile(path)
    base = load_rules()
    hierarchy = profile.get("resolved_hierarchy") or []
    body_sizes = profile["body"].get("level_sizes", [14, 13, 12, 11, 10])
    body_bold = profile["body"].get("level_bold", [True, True, False, False, False])
    default_bullets = ["auto:1.", "auto:1)", "wingdings:❑", "arial:–", "arial:•"]

    # Wingdings 글머리는 사설영역 코드(U+F0xx)로 저장돼 그대로 두면 화면에서 깨져 보인다.
    wingdings_map = {"": "wingdings:❑", "": "wingdings:■", "": "wingdings:◆", "": "wingdings:▪"}
    levels: list[dict[str, Any]] = []
    for i in range(5):
        node = hierarchy[i] if i < len(hierarchy) else {}
        bullet = node.get("bullet") or ""
        if bullet == "arabicPeriod":
            bullet_label = "auto:1."
        elif bullet == "arabicParenR":
            bullet_label = "auto:1)"
        elif bullet in ("", "none"):
            bullet_label = default_bullets[i]
        elif bullet in wingdings_map:
            bullet_label = wingdings_map[bullet]
        else:
            bullet_label = f"char:{bullet}"
        levels.append({
            "level": i + 1,
            "size": float(node.get("size_pt") or body_sizes[min(i, len(body_sizes) - 1)]),
            "bold": bool(body_bold[min(i, len(body_bold) - 1)]),
            "marL": int(round(float(node.get("margin_left_in", 0)) * 914400)),
            "indent": int(round(float(node.get("indent_in", 0)) * 914400)),
            "bullet": bullet_label,
        })

    fonts = dict(base.get("fonts") or {})
    fonts.setdefault("latin", profile["body"].get("font_latin", "Corbel"))
    fonts.setdefault("korean", profile["body"].get("font_ea", "나눔스퀘어"))
    fonts.setdefault("heading_korean", profile["title"].get("font_ea", "나눔스퀘어 ExtraBold"))
    # 새 양식이 다른 글꼴을 쓰면 그 값을 따른다.
    fonts["latin"] = profile["body"].get("font_latin", fonts["latin"])
    fonts["korean"] = profile["body"].get("font_ea", fonts["korean"])
    fonts["heading_korean"] = profile["title"].get("font_ea", fonts["heading_korean"])

    table = profile["table_type1"]
    derived = {
        "canvas": {"width_in": profile["canvas"]["width"], "height_in": profile["canvas"]["height"]},
        "fonts": fonts,
        "title": {
            "font_size": float(profile["title"].get("font_size", 24)),
            "bold": bool(profile["title"].get("bold", True)),
            "x": float(profile["title"].get("x", 0)),
            "y": float(profile["title"].get("y", 0)),
        },
        "body_levels": levels,
        "table": {
            "x_in": float(table.get("x", 2.89757)),
            "width_in": float(table.get("w", 7.519097)),
            "caption_offset_in": float((base.get("table") or {}).get("caption_offset_in", -0.367)),
            "header_fill": table.get("header_fill", "#DCE6F2"),
            "header_font_size": float(table.get("header_font_size", 10.5)),
            "body_font_size": float(table.get("body_font_size", 10)),
            "first_col_fill": (base.get("table") or {}).get("first_col_fill", "#E8E8E8"),
            "gridline": (base.get("table") or {}).get("gridline", {"width_pt": 0.75, "color": "#BFBFBF"}),
        },
        "frame": {
            "header_fill": profile["frame"].get("header_fill", "#B7D3EE"),
            "border_color": profile["frame"].get("border_color", "#7F7F7F"),
            "header_font_size": float(profile["frame"].get("header_font_size", 14)),
        },
        "footnote": {"font_size": float(profile["footnote"].get("font_size", 8))},
        "continuation": base.get("continuation") or {
            "rule": "동일 대주제가 다음 장으로 이어지면 Lv1 제목 색을 #FFFFFF 로 바꿔 숨긴다"
        },
        "cleanup": base.get("cleanup") or {"remove_empty_textbox": True, "collapse_spaces": True},
        "source": {"file": path.name, "extracted": True},
    }
    return derived


# ── 추출된 슬라이드 데이터 검사 · 자동 수정 (페이지 편집용) ──────

def _level_sizes(rules: dict[str, Any]) -> list[float]:
    levels = rules.get("body_levels") or []
    sizes = [float(x.get("size", s)) for x, s in zip(levels, [14, 13, 12, 11, 10])]
    return sizes or [14, 13, 12, 11, 10]


def validate_slides(payload: dict[str, Any]) -> dict[str, Any]:
    """추출된 슬라이드 데이터(편집 대상)를 검사한다.

    validate_pptx가 '업로드 원본'을 보는 것과 달리, 이쪽은 편집기에서 실제로
    고칠 수 있는 항목만 본다.
    """
    rules = load_rules()
    cleanup = rules.get("cleanup") or {}
    max_level = len(_level_sizes(rules)) - 1
    issues: list[dict[str, Any]] = []

    def add(category: str, slide_no: int, message: str, detail: str, field: str, index: int | None = None) -> None:
        issues.append({
            "category": category,
            "severity": CATEGORY_SEVERITY[category],
            "slide_no": slide_no,
            "message": message,
            "detail": detail,
            "auto_fixable": True,
            "field": field,
            "index": index,
        })

    for slide in payload.get("slides", []):
        no = slide["slide_number"]
        for i, item in enumerate(slide.get("body", [])):
            text = str(item.get("text", ""))
            stripped = text.strip()
            if stripped and stripped[0] in BAD_BULLET_CHARS:
                add("글머리 기호", no, f"'{stripped[0]}' → 표준 글머리로 교체",
                    f"{i + 1}번째 문단 “{stripped[:24]}”", "body", i)
            if cleanup.get("collapse_spaces", True) and re.search(r"  +", text):
                add("정리", no, "연속 공백 → 1칸으로 축소", f"{i + 1}번째 문단", "body", i)
            if int(item.get("level", 0)) > max_level:
                add("들여쓰기", no, f"{int(item['level']) + 1}단계 → {max_level + 1}단계로 정리",
                    f"{i + 1}번째 문단", "body", i)
            if not stripped:
                add("정리", no, "빈 문단 제거", f"{i + 1}번째 문단", "body", i)
        title = str(slide.get("page_title", ""))
        if title.strip() and title.strip()[0] in BAD_BULLET_CHARS:
            add("글머리 기호", no, f"제목의 '{title.strip()[0]}' 기호 제거", "페이지 제목", "page_title", None)
        if cleanup.get("collapse_spaces", True) and re.search(r"  +", title):
            add("정리", no, "제목의 연속 공백 축소", "페이지 제목", "page_title", None)

    return summarize(issues)


def preview_autofix(payload: dict[str, Any], slide_no: int | None = None) -> list[dict[str, Any]]:
    """자동 수정이 무엇을 어떻게 바꿀지 미리 계산한다 (적용하지 않음).

    화면에서 before/after를 빨강·초록으로 대조해 보여주기 위한 데이터.
    """
    import copy

    before = copy.deepcopy(payload)
    after = copy.deepcopy(payload)
    autofix_slides(after, slide_no)

    changes: list[dict[str, Any]] = []
    for src, dst in zip(before.get("slides", []), after.get("slides", [])):
        no = src["slide_number"]
        if slide_no is not None and no != slide_no:
            continue
        if src.get("page_title") != dst.get("page_title"):
            changes.append({
                "slide_no": no, "field": "제목", "index": None,
                "before": src.get("page_title", ""), "after": dst.get("page_title", ""),
            })
        src_body = src.get("body", [])
        dst_body = dst.get("body", [])
        dst_texts = [x.get("text", "") for x in dst_body]
        used = 0
        for i, item in enumerate(src_body):
            text = item.get("text", "")
            level = int(item.get("level", 0))
            if used < len(dst_body) and _matches(text, dst_texts[used]):
                new_item = dst_body[used]
                used += 1
                if text != new_item.get("text", "") or level != int(new_item.get("level", 0)):
                    changes.append({
                        "slide_no": no, "field": f"본문 {i + 1}번째 문단", "index": i,
                        "before": text, "after": new_item.get("text", ""),
                        "before_level": level, "after_level": int(new_item.get("level", 0)),
                    })
            else:
                changes.append({
                    "slide_no": no, "field": f"본문 {i + 1}번째 문단", "index": i,
                    "before": text, "after": "", "removed": True,
                })
    return changes


def _matches(before: str, after: str) -> bool:
    """자동 수정 전후 문단이 같은 문단인지 판단 (기호·공백 정리를 감안)."""
    a = re.sub(r"\s+", "", before)
    b = re.sub(r"\s+", "", after)
    if not b:
        return False
    a_clean = a[1:] if a and a[0] in BAD_BULLET_CHARS else a
    return a_clean == b or a == b or b in a_clean or a_clean in b


def autofix_slides(payload: dict[str, Any], slide_no: int | None = None) -> int:
    """추출 데이터에 자동 수정을 적용하고 고친 건수를 돌려준다."""
    rules = load_rules()
    cleanup = rules.get("cleanup") or {}
    collapse = cleanup.get("collapse_spaces", True)
    max_level = len(_level_sizes(rules)) - 1
    fixed = 0

    def clean(text: str) -> tuple[str, int]:
        n = 0
        out = str(text)
        stripped = out.strip()
        if stripped and stripped[0] in BAD_BULLET_CHARS:
            out = stripped[1:].strip()
            n += 1
        if collapse and re.search(r"  +", out):
            out = re.sub(r" {2,}", " ", out).strip()
            n += 1
        return out, n

    for slide in payload.get("slides", []):
        if slide_no is not None and slide["slide_number"] != slide_no:
            continue
        title, n = clean(slide.get("page_title", ""))
        slide["page_title"] = title
        fixed += n
        body = []
        for item in slide.get("body", []):
            text, n = clean(item.get("text", ""))
            fixed += n
            if not text.strip():
                fixed += 1
                continue
            level = int(item.get("level", 0))
            if level > max_level:
                level = max_level
                fixed += 1
            body.append({"text": text, "level": level, "bold": item.get("bold")})
        if body:
            slide["body"] = body
    return fixed
